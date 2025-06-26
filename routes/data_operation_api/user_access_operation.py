import base64
import traceback
from io import BytesIO
from flask import jsonify, Blueprint, request, Response
from openpyxl.workbook import Workbook
from sqlalchemy import func, distinct, create_engine, text

from models import db
from models.message import Message
from models.user_access import UserAccessInfo
from utils.constant import Constants
from utils.permission_utils import operation_page_permission

operation_api_bp = Blueprint('operation_api', __name__)

@operation_api_bp.route('/home_page_or_chat', methods=['GET'])
def operation_data():
    try:
        # 运维数据 首页 云集chat人员使用次数
        query_result = db.session.query(
            func.count(UserAccessInfo.user_name).label('user_count'),
            UserAccessInfo.access_platform
        ).group_by(UserAccessInfo.access_platform).all()

        stats = [{
            'platform': item.access_platform,
            'access_count': item.user_count
        } for item in query_result]
        return jsonify({'data': stats, 'status': 'success'})
    except Exception as e:
        traceback.print_exc()
        return jsonify({'data': str(e), 'status': 'fail'}),500
    finally:
        db.session.close()

@operation_api_bp.route('/operation_permission/<user_name>', methods=['POST'])
def operation_permission(user_name):
    if operation_page_permission(user_name):
        return jsonify({'data': 'success', 'status': 'success'})
    else:
        return jsonify({'data': 'fail', 'status': 'fail'})

@operation_api_bp.route('/active_users', methods=['POST'])
def active_users_operation_data():
    # 活跃用户数据
    data = request.get_json()
    start_time = data.get('start_time')
    end_time =  data.get('end_time')
    try:
        result = db.session.query(
            func.count(distinct(Message.domain_account)).label('userNum'),
            func.count(Message.message_id).label('callNum')
        ).filter(
            Message.access_time >= start_time,
            Message.access_time <= end_time
        ).first()
        ret_result = {
            "userNum" : result.userNum,
            "callNum" : result.callNum
        }

        return jsonify({'data': ret_result, 'status': 'success'})
    except Exception as e:
        traceback.print_exc()
        return jsonify({'data': str(e), 'status': 'fail'}),500
    finally:
        # 手动释放资源（如果不在 Flask 请求上下文中）
        db.session.close()

@operation_api_bp.route('/active_users_trend', methods=['POST'])
def active_users_operation_trend():
    uri = base64.b64decode(Constants.URI).decode('utf-8')
    # 活跃用户数据
    data = request.get_json()
    start_time = data.get('start_time')
    end_time =  data.get('end_time')
    engine = create_engine(uri)

    if data.get('date_type') == 'week':
        sql = week_sql_method()
    elif data.get('date_type') == 'month':
        sql = month_sql_method()
    else:
        return jsonify({'data': "date type error", 'status': 'fail'}),500
    try:
        with engine.connect() as conn:
            result = conn.execute(text(sql), {"start_time": start_time, "end_time": end_time})
            rows = result.fetchall()
        # 转成字典列表
        keys = result.keys()
        data_list = [dict(zip(keys, row)) for row in rows]

        return jsonify({'data': data_list, 'status': 'success'})
    except Exception as e:
        traceback.print_exc()
        return jsonify({'data': str(e), 'status': 'fail'}),500

@operation_api_bp.route('/export_to_excel', methods=['POST'])
def export_to_excel():
    uri = base64.b64decode(Constants.URI).decode('utf-8')
    # 人员使用次数统计sql
    user_usage_count_sql = usage_sql_method()
    week_usage_count_sql = week_usage_count_sql_method()
    month_usage_count_sql = month_usage_count_sql_method()

    try:
        # sql 查询
        with create_engine(uri).connect() as conn:  # 只创建一次连接
            # 人员统计
            user_result = conn.execute(text(user_usage_count_sql))
            user_list = [dict(zip(user_result.keys(), row)) for row in user_result]

            # 周统计
            week_result = conn.execute(text(week_usage_count_sql))
            week_list = [dict(zip(week_result.keys(), row)) for row in week_result]

            # 月统计
            month_result = conn.execute(text(month_usage_count_sql))
            month_list = [dict(zip(month_result.keys(), row)) for row in month_result]

        return export_dict_to_excel(user_list, week_list, month_list,
                                ["员工使用次数统计", "每周用户数量", "每月用户数量"])
    except Exception as e:
        traceback.print_exc()
        return jsonify({'data': str(e), 'status': 'export error'}),500

def week_usage_count_sql_method():
    return """
            SELECT 
                YEAR(week_start_date) AS year,
                WEEK(week_start_date, 1) AS week_number,
                COUNT(DISTINCT domain_account) AS user_count
            FROM (
                -- 生成完整的自然周日期序列
                SELECT 
                    DATE_SUB(access_date, INTERVAL WEEKDAY(access_date) DAY) AS week_start_date,
                    domain_account
                FROM (
                    -- 获取去重后的日期和用户组合
                    SELECT DISTINCT
                        DATE(access_time) AS access_date,
                        domain_account
                    FROM 
                        t_lobechat_messages_flask
                    WHERE 
                        access_time >= '2024-12-01'
                        AND access_time <= CURDATE()
                ) AS distinct_dates
            ) AS weekly_data
            GROUP BY 
                week_start_date
    """

def month_usage_count_sql_method():
    return """
            SELECT 
                YEAR(access_date) AS year,
                MONTH(access_date) AS month_number,
                COUNT(DISTINCT domain_account) AS user_count
            FROM (
                SELECT DISTINCT
                    DATE(access_time) AS access_date,
                    domain_account
                FROM 
                    t_lobechat_messages_flask
                WHERE 
                    access_time >= '2024-12-01'
                    AND access_time <= CURDATE()
            ) AS distinct_dates
            GROUP BY 
                YEAR(access_date),
                MONTH(access_date)
            ORDER BY 
                year, month_number;
    """


def export_dict_to_excel(data_dict, week_list, month_list, sheet_names):
    # 参数校验
    if sheet_names is None:
        sheet_names = ["员工使用次数统计", "每周用户数量", "每月用户数量"]
    elif len(sheet_names) != 3:
        raise ValueError("sheet_names 必须包含3个名称")

    wb = Workbook()

    # 删除默认创建的Sheet（如果有）
    if "Sheet" in wb.sheetnames:
        wb.remove(wb["Sheet"])

    def _write_sheet(worksheet, data):
        """通用写入Sheet的函数"""
        if data and len(data) > 0:  # 检查数据存在且不为空
            # 写入表头
            headers = list(data[0].keys())
            for col_idx, header in enumerate(headers, start=1):
                worksheet.cell(row=1, column=col_idx, value=header)

            # 写入数据
            for row_idx, row_data in enumerate(data, start=2):
                for col_idx, header in enumerate(headers, start=1):
                    worksheet.cell(row=row_idx, column=col_idx, value=row_data.get(header, ""))
        else:
            worksheet.append(["无数据"])

    # 为每个数据集创建Sheet并写入数据
    for i, (data, sheet_name) in enumerate(zip(
            [data_dict, week_list, month_list],
            sheet_names
    )):
        ws = wb.create_sheet(title=sheet_name, index=i)
        _write_sheet(ws, data)

    # 保存到内存
    excel_data = BytesIO()
    wb.save(excel_data)
    excel_data.seek(0)

    return Response(
        excel_data.getvalue(),
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=multi_sheet_report.xlsx"}
    )

def usage_sql_method():
    return """
            SELECT 
                m.domain_account AS 工号,
                u.chn_name AS 姓名,
                COUNT(*) AS 使用次数,
                u.departmentL1,
                u.departmentL2,
                u.departmentL3,
                u.departmentL4,
                u.departmentL5,
                u.departmentL6
            FROM 
                t_lobechat_messages_flask m
            LEFT JOIN 
                t_lobechat_user_info u ON  u.user_name COLLATE utf8mb4_general_ci  = m.domain_account 
             WHERE 
                m.domain_account IS NOT NULL
            GROUP BY 
                m.domain_account
    """


def week_sql_method():
    week_sql = """
          SELECT 
            YEAR(week_start_date) AS year,
            WEEK(week_start_date, 1) AS time_number,
            COUNT(DISTINCT domain_account) AS user_count
        FROM (
            -- 生成完整的自然周日期序列
            SELECT 
                DATE_SUB(access_date, INTERVAL WEEKDAY(access_date) DAY) AS week_start_date,
                domain_account
            FROM (
                -- 获取去重后的日期和用户组合
                SELECT DISTINCT
                    DATE(access_time) AS access_date,
                    domain_account
                FROM 
                    t_lobechat_messages_flask
                WHERE 
                    access_time >= DATE_SUB(:start_time, INTERVAL WEEKDAY(:start_time) DAY)  -- 对齐周一开始
                    AND access_time <= DATE_ADD(DATE_SUB(:end_time, INTERVAL WEEKDAY(:end_time) DAY), INTERVAL 6 DAY)  -- 对齐周日结束
            ) AS distinct_dates
        ) AS weekly_data
        GROUP BY 
            week_start_date
    """
    return week_sql


def month_sql_method():
    month_sql = """
            SELECT 
                YEAR(month_start_date) AS year,
                MONTH(month_start_date) AS time_number,
                COUNT(DISTINCT domain_account) AS user_count
            FROM (
                SELECT 
                    DATE_FORMAT(access_date, '%Y-%m-01') AS month_start_date,  -- 强制对齐到月初
                    domain_account
                FROM (
                    SELECT DISTINCT
                        DATE(access_time) AS access_date,
                        domain_account
                    FROM 
                        t_lobechat_messages_flask
                    WHERE 
                        access_time >= DATE_FORMAT(:start_time, '%Y-%m-01')  -- 开始时间对齐当月1号
                        AND access_time <= LAST_DAY(:end_time)               -- 结束时间对齐当月最后一天
                ) AS distinct_dates
            ) AS monthly_data
            GROUP BY 
                YEAR(month_start_date),
                MONTH(month_start_date)
            ORDER BY
                year, time_number;
        """
    return month_sql
